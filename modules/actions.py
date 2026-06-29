
# Action Items — editable tracker with priority briefing and Action Tracker export.

import io
import numpy as np
import pandas as pd
import streamlit as st
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import (
    SECTORS, ENGAGEMENT_TYPES, ACTION_STATUSES, PROGRESS_OPTIONS,
    ESCALATION_FLAGS, TASK_PRIORITIES, DEPARTMENTS, MISA_GREEN, MISA_GOLD,
    STATUS_COLORS, PRIORITY_COLORS,
)
from config.translations import t
from modules.persistence import save_session

_GREEN = "#1B5C3F"
_GOLD  = "#C9974A"
_RED   = "#DC2626"
_AMBER = "#D97706"
_BLUE  = "#1D4ED8"
_GRAY  = "#6B7280"

_STATUS_BG = {
    "Completed":   "#D1FAE5", "In Progress": "#FEF3C7", "Inprogress":  "#FEF3C7",
    "Not Started": "#F3F4F6", "Blocked":     "#FEE2E2", "Cancelled":   "#F3F4F6",
}
_STATUS_FG = {
    "Completed":   "#065F46", "In Progress": "#92400E", "Inprogress":  "#92400E",
    "Not Started": "#374151", "Blocked":     "#991B1B", "Cancelled":   "#6B7280",
}
_PRIO_COLOR = {"High": _RED, "Medium": _AMBER, "Low": _GRAY, "Very High": _RED}

# ── Public: exported for dashboard ────────────────────────────────────────────

def render_action_summary_widget(dfs: dict, lang: str):
    """Compact 'What needs attention' panel for dashboard embedding."""
    actions = dfs.get("Action Items", pd.DataFrame())
    if actions.empty:
        return
    today  = date.today()
    urgent = _classify_actions(actions, today)

    n_blocked  = len(urgent["blocked"])
    n_overdue  = len(urgent["overdue"])
    n_today    = len(urgent["due_today"])
    n_soon     = len(urgent["due_soon"])
    total_flag = n_blocked + n_overdue + n_today

    if total_flag == 0 and n_soon == 0:
        return

    header_color = "linear-gradient(135deg,#1a0505,#7f1d1d)" if total_flag > 0 \
                   else f"linear-gradient(135deg,#071a0f,{_GREEN})"

    st.markdown(
        f'<div style="background:{header_color};border-radius:12px;'
        f'padding:14px 20px;margin-bottom:10px;">'
        f'<div style="color:{_GOLD};font-size:9px;font-weight:700;'
        f'letter-spacing:1px;text-transform:uppercase;margin-bottom:8px;">'
        f'Action Items — Attention Required</div>'
        f'<div style="display:flex;gap:14px;flex-wrap:wrap;">',
        unsafe_allow_html=True,
    )
    for label, count, color in [
        ("Blocked",   n_blocked, "#FCA5A5"),
        ("Overdue",   n_overdue, "#FCA5A5"),
        ("Due Today", n_today,   "#FCD34D"),
        ("Due Soon",  n_soon,    "rgba(255,255,255,0.6)"),
    ]:
        if count == 0:
            continue
        st.markdown(
            f'<div style="text-align:center;">'
            f'<div style="color:{color};font-size:22px;font-weight:700;">{count}</div>'
            f'<div style="color:rgba(255,255,255,0.55);font-size:9px;">{label}</div>'
            f'</div>',
            unsafe_allow_html=True,
        )
    st.markdown('</div></div>', unsafe_allow_html=True)

    # Top-4 critical items — white card so text is always visible
    top_items = (urgent["blocked"] + urgent["overdue"] + urgent["due_today"])[:4]
    for item in top_items:
        sc = _RED if item["urgency"] in ("blocked", "overdue") else _AMBER
        st.markdown(
            f'<div style="display:flex;align-items:flex-start;gap:8px;'
            f'padding:6px 10px;border-radius:6px;margin-bottom:4px;'
            f'background:#fff;border-left:3px solid {sc};'
            f'box-shadow:0 1px 3px rgba(0,0,0,0.07);">'
            f'<div style="flex:1;min-width:0;">'
            f'<div style="display:flex;align-items:center;gap:6px;margin-bottom:2px;">'
            f'<span style="color:{_GREEN};font-size:10px;font-weight:700;">{item["company"]}</span>'
            f'<span style="background:{sc};color:#fff;padding:1px 5px;border-radius:3px;'
            f'font-size:8px;font-weight:700;">{item["tag"]}</span>'
            f'</div>'
            f'<div style="color:#1F2937;font-size:10px;line-height:1.4;">{item["desc"][:80]}</div>'
            f'</div></div>',
            unsafe_allow_html=True,
        )


# ── Main render ────────────────────────────────────────────────────────────────

def render(dfs: dict, lang: str):
    actions   = dfs.get("Action Items",    pd.DataFrame())
    investors = dfs.get("Investor Master", pd.DataFrame())
    today     = date.today()

    # ── Header row ────────────────────────────────────────────────────────────
    hcol, ecol = st.columns([5, 1])
    hcol.markdown(
        f"<h2 style='color:{_GREEN};margin-bottom:2px;'>Action Items</h2>",
        unsafe_allow_html=True,
    )

    # ── Priority briefing panel ───────────────────────────────────────────────
    if not actions.empty:
        _render_priority_panel(actions, today)

    # ── Add action ────────────────────────────────────────────────────────────
    with st.expander("➕ Add Action Item", expanded=False):
        _add_action_form(dfs, investors, lang)

    # ── Filters ───────────────────────────────────────────────────────────────
    fc1, fc2, fc3, fc4 = st.columns([2.5, 1.5, 1.5, 1.5])
    with fc1:
        q = fc1.text_input("Search", placeholder="Search action or company…",
                           label_visibility="collapsed", key="act_q")
    companies   = sorted(actions["Company Name"].dropna().unique().tolist()) \
                  if not actions.empty and "Company Name" in actions.columns else []
    filter_co   = fc2.multiselect("Company",  companies, label_visibility="collapsed",
                                  placeholder="All companies", key="act_co")
    filter_stat = fc3.multiselect("Status",   ACTION_STATUSES, label_visibility="collapsed",
                                  placeholder="All statuses", key="act_st")
    filter_pri  = fc4.multiselect("Priority", ["High", "Medium", "Low"],
                                  label_visibility="collapsed",
                                  placeholder="All priorities", key="act_pr")

    filtered = actions.copy()
    if not filtered.empty:
        if q:
            mask = (
                filtered.get("Action Description", pd.Series(dtype=str)).fillna("").str.contains(q, case=False, na=False)
                | filtered.get("Company Name", pd.Series(dtype=str)).fillna("").str.contains(q, case=False, na=False)
            )
            filtered = filtered[mask]
        if filter_co   and "Company Name" in filtered.columns:
            filtered = filtered[filtered["Company Name"].isin(filter_co)]
        if filter_stat and "Status"       in filtered.columns:
            filtered = filtered[filtered["Status"].isin(filter_stat)]
        if filter_pri  and "Priority"     in filtered.columns:
            filtered = filtered[filtered["Priority"].isin(filter_pri)]

    # ── Status summary strip ──────────────────────────────────────────────────
    if not filtered.empty and "Status" in filtered.columns:
        _render_status_strip(filtered)

    if filtered.empty:
        st.info("No action items match the current filters.")
        return

    # ── Sort: pending by due date first ───────────────────────────────────────
    display = filtered.copy()
    if "Due Date" in display.columns:
        display["_due"] = pd.to_datetime(display["Due Date"], errors="coerce")
        pending = display[~display.get("Status", pd.Series(dtype=str)).isin(["Completed", "Cancelled"])]
        done    = display[ display.get("Status", pd.Series(dtype=str)).isin(["Completed", "Cancelled"])]
        pending = pending.sort_values("_due", na_position="last")
        display = pd.concat([pending, done], ignore_index=True).drop(columns=["_due"], errors="ignore")

    # Flag overdue rows
    if "Due Date" in display.columns and "Status" in display.columns:
        due_dt      = pd.to_datetime(display["Due Date"], errors="coerce")
        is_pending  = ~display["Status"].isin(["Completed", "Cancelled"])
        display["⚠"] = [
            "🔴 Overdue" if (pd.notna(d) and d.date() < today and is_pending.iloc[i]) else ""
            for i, d in enumerate(due_dt)
        ]

    edit_cols = [c for c in [
        "⚠", "Action ID", "Company Name", "Action Description",
        "Assigned To", "Type of Engagement", "Priority",
        "Status", "Progress", "Due Date", "Escalation Flag", "Remarks", "AM Input",
    ] if c in display.columns]

    col_cfg = {
        "⚠":                  st.column_config.TextColumn("⚠", width="small", disabled=True),
        "Action ID":          st.column_config.TextColumn("ID", width="small", disabled=True),
        "Company Name":       st.column_config.TextColumn("Company", disabled=True),
        "Action Description": st.column_config.TextColumn("Action Item", width="large"),
        "Assigned To":        st.column_config.TextColumn("Owner"),
        "Due Date":           st.column_config.DateColumn("Due Date"),
        "Status":             st.column_config.SelectboxColumn("Status",   options=ACTION_STATUSES),
        "Priority":           st.column_config.SelectboxColumn("Priority", options=["High", "Medium", "Low", "Very High"]),
        "Progress":           st.column_config.SelectboxColumn("Progress", options=PROGRESS_OPTIONS),
        "Escalation Flag":    st.column_config.SelectboxColumn("Escalation", options=ESCALATION_FLAGS),
        "Type of Engagement": st.column_config.SelectboxColumn("Type",    options=ENGAGEMENT_TYPES),
        "Remarks":            st.column_config.TextColumn("Remarks"),
        "AM Input":           st.column_config.TextColumn("AM Input"),
    }

    st.markdown(
        '<div style="font-size:11px;color:#9CA3AF;margin-bottom:4px;">'
        '✏️ Click any cell to edit. Changes are saved when you click <strong>Save Changes</strong>.'
        '</div>',
        unsafe_allow_html=True,
    )

    edited_df = st.data_editor(
        display[edit_cols],
        key="actions_editor",
        use_container_width=True,
        hide_index=True,
        num_rows="fixed",
        column_config=col_cfg,
    )

    # ── Action buttons ────────────────────────────────────────────────────────
    sa_col, ex_col, _ = st.columns([1.4, 1.4, 3])

    if sa_col.button("💾 Save Changes", type="primary", use_container_width=True):
        n_saved = _merge_edits_back(dfs, edited_df, display[edit_cols])
        if n_saved:
            st.success(f"✅ Saved {n_saved} change{'s' if n_saved!=1 else ''}")
            st.rerun()
        else:
            st.info("No changes detected.")

    if ex_col.button("📥 Export Action Tracker", use_container_width=True):
        xl = _export_tracker_excel(dfs)
        fname = f"ActionTracker_{date.today().strftime('%Y-%m-%d')}.xlsx"
        st.download_button(
            "⬇ Download Action Tracker Excel",
            data=xl, file_name=fname,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True, key="act_dl",
        )

    st.caption(f"{len(display)} action items")


# ── Priority briefing panel ────────────────────────────────────────────────────

def _classify_actions(actions: pd.DataFrame, today) -> dict:
    """Return actions bucketed by urgency."""
    buckets: dict[str, list] = {
        "blocked": [], "overdue": [], "due_today": [], "due_soon": [],
    }
    if actions.empty:
        return buckets

    due_col = pd.to_datetime(actions.get("Due Date", pd.Series(dtype=str)), errors="coerce") \
              if "Due Date" in actions.columns else pd.Series([pd.NaT] * len(actions))
    for i, row in actions.iterrows():
        status = str(row.get("Status", "")).strip()
        if status in ("Completed", "Cancelled"):
            continue
        due = due_col.iloc[i] if i < len(due_col) else pd.NaT
        item = {
            "company": str(row.get("Company Name", "?")),
            "desc":    str(row.get("Action Description", ""))[:90],
            "owner":   str(row.get("Assigned To", "—")),
            "priority":str(row.get("Priority", "Medium")),
            "due":     due.date() if pd.notna(due) else None,
        }
        if status == "Blocked":
            item["urgency"] = "blocked"
            item["tag"] = "BLOCKED"
            buckets["blocked"].append(item)
        elif pd.notna(due):
            d = due.date()
            if d < today:
                item["urgency"] = "overdue"
                diff = (today - d).days
                item["tag"] = f"OVERDUE {diff}d"
                buckets["overdue"].append(item)
            elif d == today:
                item["urgency"] = "due_today"
                item["tag"] = "DUE TODAY"
                buckets["due_today"].append(item)
            elif d <= today + timedelta(days=7):
                item["urgency"] = "due_soon"
                days_left = (d - today).days
                item["tag"] = f"DUE IN {days_left}d"
                buckets["due_soon"].append(item)
    return buckets


def _render_priority_panel(actions: pd.DataFrame, today):
    buckets = _classify_actions(actions, today)
    n_blocked = len(buckets["blocked"])
    n_overdue = len(buckets["overdue"])
    n_today   = len(buckets["due_today"])
    n_soon    = len(buckets["due_soon"])
    total_critical = n_blocked + n_overdue + n_today

    if total_critical == 0 and n_soon == 0:
        st.markdown(
            f'<div style="background:#f0fdf4;border:1px solid #86efac;border-radius:8px;'
            f'padding:10px 16px;margin-bottom:12px;color:#065F46;font-size:13px;">'
            f'✅ No overdue, blocked, or urgent actions. All items are on track.</div>',
            unsafe_allow_html=True,
        )
        return

    hdr_bg = "linear-gradient(135deg,#1a0505 0%,#7f1d1d 100%)" if total_critical > 0 \
             else f"linear-gradient(135deg,#071a0f 0%,{_GREEN} 80%)"

    # Header banner
    st.markdown(
        f'<div style="background:{hdr_bg};border-radius:12px 12px 0 0;padding:14px 20px;">'
        f'<div style="display:flex;align-items:center;justify-content:space-between;'
        f'flex-wrap:wrap;gap:10px;">'
        f'<div>'
        f'<div style="color:{_GOLD};font-size:9px;font-weight:700;letter-spacing:1px;'
        f'text-transform:uppercase;margin-bottom:4px;">What needs your attention today</div>'
        f'<div style="color:#fff;font-size:18px;font-weight:700;">'
        f'{total_critical} item{"s" if total_critical!=1 else ""} require immediate action</div>'
        f'</div>'
        f'<div style="display:flex;gap:14px;">'
        + "".join(
            f'<div style="text-align:center;">'
            f'<div style="color:{c};font-size:24px;font-weight:700;">{n}</div>'
            f'<div style="color:rgba(255,255,255,0.55);font-size:9px;text-transform:uppercase;">{lbl}</div>'
            f'</div>'
            for lbl, n, c in [
                ("Blocked",   n_blocked, "#FCA5A5"),
                ("Overdue",   n_overdue, "#FCA5A5"),
                ("Due Today", n_today,   "#FCD34D"),
                ("Due Soon",  n_soon,    "rgba(255,255,255,0.6)"),
            ] if n > 0
        )
        + f'</div></div></div>',
        unsafe_allow_html=True,
    )

    # Priority item cards — light background so dark text is always visible
    all_urgent = (
        [(i, "blocked",   _RED)   for i in buckets["blocked"]]
        + [(i, "overdue",   _RED)   for i in buckets["overdue"]]
        + [(i, "due_today", _AMBER) for i in buckets["due_today"]]
        + [(i, "due_soon",  _GRAY)  for i in buckets["due_soon"]]
    )[:10]

    st.markdown("<div style='margin-bottom:16px;'>", unsafe_allow_html=True)
    cols = st.columns(2)
    for idx, (item, urgency, border_c) in enumerate(all_urgent):
        prio_badge = ""
        if item["priority"] in ("High", "Very High"):
            pc = _PRIO_COLOR.get(item["priority"], _RED)
            prio_badge = (f'<span style="background:{pc};color:#fff;padding:1px 5px;'
                          f'border-radius:4px;font-size:8px;margin-left:4px;">'
                          f'{item["priority"]}</span>')
        tag_bg = _RED if urgency in ("blocked", "overdue") else \
                 _AMBER if urgency == "due_today" else _GRAY
        due_str = f'  ·  📅 {item["due"].strftime("%d %b")}' if item["due"] else ""
        with cols[idx % 2]:
            st.markdown(
                f'<div style="background:#fff;border:1px solid #E5E7EB;border-left:3px solid {border_c};'
                f'border-radius:7px;padding:8px 12px;margin-bottom:8px;'
                f'box-shadow:0 1px 3px rgba(0,0,0,0.06);">'
                f'<div style="display:flex;align-items:center;gap:6px;margin-bottom:4px;">'
                f'<span style="color:{_GREEN};font-size:11px;font-weight:700;">{item["company"]}</span>'
                f'<span style="background:{tag_bg};color:#fff;padding:1px 6px;border-radius:4px;'
                f'font-size:8px;font-weight:700;">{item["tag"]}</span>'
                f'{prio_badge}'
                f'</div>'
                f'<div style="color:#1F2937;font-size:11px;line-height:1.4;margin-bottom:3px;">'
                f'{item["desc"]}</div>'
                f'<div style="color:#6B7280;font-size:9px;">'
                f'👤 {item["owner"]}{due_str}</div>'
                f'</div>',
                unsafe_allow_html=True,
            )
    st.markdown("</div>", unsafe_allow_html=True)


# ── Status summary strip ──────────────────────────────────────────────────────

def _render_status_strip(df: pd.DataFrame):
    counts = df["Status"].value_counts()
    cols   = st.columns(len(ACTION_STATUSES))
    for i, status in enumerate(ACTION_STATUSES):
        count = counts.get(status, 0)
        color = STATUS_COLORS.get(status, "#9B9B9B")
        cols[i].markdown(
            f"<div style='background:{color};color:#fff;padding:8px 4px;"
            f"border-radius:6px;text-align:center;font-size:12px;'>"
            f"<b>{count}</b><br>{status}</div>",
            unsafe_allow_html=True,
        )
    st.markdown("")


# ── Edit/save helpers ─────────────────────────────────────────────────────────

def _merge_edits_back(dfs: dict, edited: pd.DataFrame, original: pd.DataFrame) -> int:
    """Merge changed rows from `edited` back into dfs['Action Items']. Returns count changed."""
    editable_cols = [
        "Action Description", "Assigned To", "Type of Engagement",
        "Priority", "Status", "Progress", "Due Date", "Escalation Flag", "Remarks", "AM Input",
    ]
    check_cols = [c for c in editable_cols if c in edited.columns and c in original.columns]
    if not check_cols:
        return 0

    full = dfs.get("Action Items", pd.DataFrame()).copy()
    if full.empty or "Action ID" not in full.columns:
        return 0

    n_changed = 0
    orig_reset = original.reset_index(drop=True)
    edit_reset = edited.reset_index(drop=True)

    for i in range(min(len(orig_reset), len(edit_reset))):
        orig_row = orig_reset.iloc[i]
        edit_row = edit_reset.iloc[i]
        act_id   = orig_row.get("Action ID")
        if not act_id:
            continue
        changed = any(
            str(orig_row.get(c, "")) != str(edit_row.get(c, ""))
            for c in check_cols
        )
        if changed:
            mask = full["Action ID"] == act_id
            for c in check_cols:
                if c in full.columns:
                    val = edit_row.get(c)
                    if "datetime" in str(full[c].dtype):
                        ts = pd.to_datetime(val, errors="coerce")
                        full[c] = full[c].mask(mask, ts)
                    else:
                        full.loc[mask, c] = val
            now_ts = pd.to_datetime(datetime.now())
            full["Last Updated"] = full["Last Updated"].mask(mask, now_ts)
            n_changed += 1

    if n_changed:
        dfs["Action Items"] = full
        save_session(dfs)
    return n_changed


# ── Export: Action Tracker Excel format ──────────────────────────────────────

def _export_tracker_excel(dfs: dict) -> bytes:
    """
    Build an Excel workbook in the Action Tracker format:
    one 'Action Items [Company]' sheet per company,
    header block matching the master template, data table at row 20.
    """
    actions   = dfs.get("Action Items",        pd.DataFrame())
    investors = dfs.get("Investor Master",      pd.DataFrame())
    opps      = dfs.get("Opportunity Pipeline", pd.DataFrame())

    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    companies = sorted(actions["Company Name"].dropna().unique().tolist()) \
                if not actions.empty and "Company Name" in actions.columns else []

    if not companies:
        ws = wb.create_sheet("Action Items")
        ws["G19"] = "No action items found"
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    for company in companies:
        sname   = f"Action Items {company}"[:31]
        ws      = wb.create_sheet(sname)
        co_acts = actions[actions["Company Name"] == company] \
                  if "Company Name" in actions.columns else pd.DataFrame()
        inv_row = investors[investors["Company Name"] == company].iloc[0] \
                  if not investors.empty and "Company Name" in investors.columns \
                  and company in investors["Company Name"].values else pd.Series()
        _write_tracker_sheet(ws, company, inv_row, co_acts, opps)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_tracker_sheet(ws, company: str, inv_row, actions: pd.DataFrame,
                         opps: pd.DataFrame = None):
    """Write one company sheet matching the master ActionTracker template."""
    green_fill  = PatternFill("solid", fgColor="1B5C3F")
    gold_fill   = PatternFill("solid", fgColor="C9974A")
    white_font  = Font(color="FFFFFF", bold=True, size=10)
    gold_font   = Font(color="FFFFFF", size=10)
    dark_font   = Font(size=10)
    bold_font   = Font(bold=True, size=10)
    center_al   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_al     = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin        = Side(style="thin")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    today       = date.today()

    inv = inv_row if (hasattr(inv_row, "empty") and not inv_row.empty) else pd.Series()

    # ── Company name banner (row 10, G-P merged) ──────────────────────────────
    ws.merge_cells("G10:P10")
    c = ws["G10"]
    c.value     = company.upper()
    c.fill      = green_fill
    c.font      = Font(color="FFFFFF", bold=True, size=14)
    c.alignment = center_al
    ws.row_dimensions[10].height = 30

    # ── Company card rows 12-16 ───────────────────────────────────────────────
    am  = str(inv.get("Account Manager",      "") or "")
    rm  = str(inv.get("Relationship Manager", "") or "")
    out = str(inv.get("Outreach Manager",     "") or "") or "Majed Alsaadi"

    rep_name  = str(inv.get("Company Rep",   "") or "")
    rep_title = str(inv.get("Rep Position",  "") or "")
    website   = str(inv.get("Website",       "") or "")
    email     = str(inv.get("Rep Email",     "") or "")
    phone     = str(inv.get("Rep Phone",     "") or "")
    country   = str(inv.get("Country",       "") or "")
    sector    = str(inv.get("Sector",        "") or "")
    inv_type  = str(inv.get("Investor Tier", "") or "Investor")
    last_upd  = today.strftime("%d %b %Y")

    for r in range(12, 17):
        ws.row_dimensions[r].height = 18

    # Left block: Outreach / AM / RM (cols K-L = 11-12)
    left_items = [("Outreach", out), ("AM", am), ("RM", rm), ("", ""), ("", "")]
    for i, (lbl, val) in enumerate(left_items):
        lc = ws.cell(row=12 + i, column=11, value=lbl)
        lc.fill = green_fill; lc.font = white_font; lc.alignment = center_al
        vc = ws.cell(row=12 + i, column=12, value=val)
        vc.font = bold_font; vc.alignment = left_al

    # Opportunity block: up to 5 opps (cols H-I = 8-9), rows 12-16
    co_opps = pd.DataFrame()
    if opps is not None and not opps.empty and "Company Name" in opps.columns:
        co_opps = opps[opps["Company Name"] == company]
    for i in range(5):
        r = 12 + i
        if i < len(co_opps):
            orow     = co_opps.iloc[i]
            opp_name = str(orow.get("Opportunity Name", "") or "")
            opp_val  = orow.get("Est. Investment Value (SAR)", None)
            nc = ws.cell(row=r, column=8, value=opp_name)
            nc.font = dark_font; nc.alignment = left_al
            if opp_val is not None and str(opp_val) not in ("", "nan"):
                vc2 = ws.cell(row=r, column=9, value=opp_val)
                vc2.font = dark_font; vc2.alignment = center_al

    # Right block 1: Type / Rep / Position / Website / Last Updated (cols M-N = 13-14)
    right1 = [("Type", inv_type), ("Rep", rep_name), ("Position", rep_title),
              ("Website", website), ("Last Updated", last_upd)]
    for i, (lbl, val) in enumerate(right1):
        lc = ws.cell(row=12 + i, column=13, value=lbl)
        lc.fill = green_fill; lc.font = white_font; lc.alignment = center_al
        vc = ws.cell(row=12 + i, column=14, value=val)
        vc.fill = gold_fill; vc.font = gold_font; vc.alignment = left_al

    # Right block 2: Company Name / Email / Phone / Country / Sector (cols O-P = 15-16)
    right2 = [("Company Name", company), ("Email", email), ("Phone", phone),
              ("Country", country), ("Sector", sector)]
    for i, (lbl, val) in enumerate(right2):
        lc = ws.cell(row=12 + i, column=15, value=lbl)
        lc.fill = green_fill; lc.font = white_font; lc.alignment = center_al
        vc = ws.cell(row=12 + i, column=16, value=val)
        vc.fill = gold_fill; vc.font = gold_font; vc.alignment = left_al

    # ── Column headers at row 19 ──────────────────────────────────────────────
    HEADERS = [
        ("G", "ID"), ("H", "Action Item"), ("I", "Assigned to"),
        ("J", "Type of Engagement"), ("K", "Start Date"), ("L", "Due Date"),
        ("M", "Priority"), ("N", "Progress"), ("O", "Remarks"), ("P", "AM Input"),
    ]
    for col_letter, hdr_text in HEADERS:
        cell = ws[f"{col_letter}19"]
        cell.value     = hdr_text
        cell.fill      = green_fill
        cell.font      = white_font
        cell.alignment = center_al
        cell.border    = border
    ws.row_dimensions[19].height = 22

    # ── Data rows from row 20 ─────────────────────────────────────────────────
    STATUS_FILLS_MAP = {
        "Completed":   "E2EFDA", "In Progress": "FFF2CC", "Inprogress": "FFF2CC",
        "Not Started": "F2F2F2", "Blocked":     "FCE4D6", "Cancelled":  "EDEDED",
    }
    if not actions.empty:
        df = actions.copy()
        if "Due Date" in df.columns:
            df["_due"] = pd.to_datetime(df["Due Date"], errors="coerce")
            pend = df[~df.get("Status", pd.Series(dtype=str)).isin(["Completed", "Cancelled"])]
            done = df[ df.get("Status", pd.Series(dtype=str)).isin(["Completed", "Cancelled"])]
            df   = pd.concat([pend.sort_values("_due"), done], ignore_index=True).drop(columns=["_due"], errors="ignore")

        for seq, (_, row) in enumerate(df.iterrows(), start=1):
            r      = 19 + seq   # row 20, 21, …
            status = str(row.get("Status", "Not Started") or "Not Started")
            due_raw = row.get("Due Date")
            due_val = None
            try:
                dv = pd.to_datetime(due_raw)
                due_val = dv.date() if pd.notna(dv) else None
            except Exception:
                pass

            is_overdue = due_val and due_val < today and status not in ("Completed", "Cancelled")
            row_fill   = PatternFill("solid", fgColor="FCE4D6") if is_overdue else \
                         PatternFill("solid", fgColor="FFFFFF" if seq % 2 == 0 else "F7F7F2")

            prog_raw = row.get("Progress", 0)
            try:
                prog_pct = float(str(prog_raw).replace("%", "")) if prog_raw not in (None, "", "nan") else 0
                if prog_pct > 1.0:
                    prog_pct /= 100.0
            except Exception:
                prog_pct = 0.0

            values = {
                "G": seq,
                "H": str(row.get("Action Description", "") or ""),
                "I": str(row.get("Assigned To",        "") or ""),
                "J": str(row.get("Type of Engagement", "") or ""),
                "K": row.get("Start Date") if row.get("Start Date") and pd.notna(row.get("Start Date")) else None,
                "L": due_val,
                "M": str(row.get("Priority", "") or ""),
                "N": f"{int(prog_pct * 100)}%",
                "O": str(row.get("Remarks",  "") or ""),
                "P": str(row.get("AM Input", "") or ""),
            }
            for col_letter, val in values.items():
                cell = ws[f"{col_letter}{r}"]
                cell.value     = val
                cell.fill      = row_fill
                cell.font      = dark_font
                cell.border    = border
                cell.alignment = center_al if col_letter in ("G", "M", "N") else left_al
            ws.row_dimensions[r].height = 40

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["G"].width = 6
    ws.column_dimensions["H"].width = 52
    ws.column_dimensions["I"].width = 18
    ws.column_dimensions["J"].width = 20
    ws.column_dimensions["K"].width = 12
    ws.column_dimensions["L"].width = 12
    ws.column_dimensions["M"].width = 11
    ws.column_dimensions["N"].width = 10
    ws.column_dimensions["O"].width = 30
    ws.column_dimensions["P"].width = 30
    ws.freeze_panes = "G20"


# ── Add action form ────────────────────────────────────────────────────────────

def _add_action_form(dfs: dict, investors: pd.DataFrame, lang: str):
    company_options = [""] + (
        sorted(investors["Company Name"].dropna().unique().tolist())
        if not investors.empty and "Company Name" in investors.columns else []
    )
    with st.form("add_action_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        company = c1.selectbox(t("company_name", lang), company_options)
        action  = c2.text_input(t("action_description", lang))
        c3, c4 = st.columns(2)
        eng_type = c3.selectbox(t("engagement_type", lang), ENGAGEMENT_TYPES)
        assigned  = c4.text_input(t("assigned_to", lang))
        c5, c6 = st.columns(2)
        priority  = c5.selectbox(t("priority", lang), ["High", "Medium", "Low"])
        progress  = c6.selectbox(t("progress", lang), PROGRESS_OPTIONS)
        c7, c8 = st.columns(2)
        start_d   = c7.date_input(t("start_date", lang), value=date.today())
        due_d     = c8.date_input(t("due_date", lang), value=None)
        c9, c10 = st.columns(2)
        escalation = c9.selectbox(t("escalation_flag", lang), ESCALATION_FLAGS)
        remarks    = c10.text_input(t("remarks", lang))
        am_input   = st.text_area("AM Input")

        if st.form_submit_button(t("add_action", lang), use_container_width=True):
            if not company or not action:
                st.warning("Company and action description are required.")
                return
            actions = dfs.get("Action Items", pd.DataFrame())
            inv_id  = _get_investor_id(investors, company)
            new_id  = _next_id(actions, "Action ID", "ACT")
            new_row = {
                "Action ID":          new_id,
                "Investor ID":        inv_id,
                "Company Name":       company,
                "Action Description": action,
                "Assigned To":        assigned,
                "Type of Engagement": eng_type,
                "Start Date":         start_d,
                "Due Date":           due_d,
                "Priority":           priority,
                "Progress":           progress,
                "Status":             "Not Started",
                "Escalation Flag":    escalation,
                "Remarks":            remarks,
                "AM Input":           am_input,
                "Last Updated":       np.datetime64(datetime.now()),
            }
            dfs["Action Items"] = pd.concat(
                [actions, pd.DataFrame([new_row])], ignore_index=True
            )
            save_session(dfs)
            st.success(f"✅ Action {new_id} added for {company}")
            st.rerun()


# ── Helpers ───────────────────────────────────────────────────────────────────

def _get_investor_id(investors: pd.DataFrame, company: str) -> str:
    if investors.empty or "Company Name" not in investors.columns:
        return ""
    m = investors[investors["Company Name"] == company]
    return str(m.iloc[0].get("Investor ID", "")) if not m.empty else ""


def _next_id(df: pd.DataFrame, id_col: str, prefix: str) -> str:
    if df.empty or id_col not in df.columns:
        return f"{prefix}-001"
    nums = []
    for v in df[id_col].dropna():
        try:
            nums.append(int(str(v).split("-")[-1]))
        except ValueError:
            pass
    return f"{prefix}-{(max(nums)+1 if nums else 1):03d}"

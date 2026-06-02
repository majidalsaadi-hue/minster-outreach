
# Auto-save and auto-load for local session persistence.
# Saves to crm_data.xlsx in the project root folder.

from pathlib import Path
from datetime import datetime
import pandas as pd

from data.schema import SCHEMA

SAVE_PATH = Path(__file__).parent.parent / "crm_data.xlsx"

# Written to session state key "save_status": "ok" | "error: <msg>"
_STATUS_KEY = "save_status"
_TIME_KEY   = "save_time"


def save_session(dfs: dict) -> bool:
    """
    Write all current DataFrames to crm_data.xlsx.
    Returns True on success, False on failure.
    Stores status in Streamlit session state for display.
    """
    import streamlit as st

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        wb.remove(wb.active)

        for sheet_name, df in dfs.items():
            if df is None:
                continue
            ws = wb.create_sheet(str(sheet_name))
            cols = list(df.columns) if len(df.columns) > 0 else []
            if not cols:
                continue
            ws.append(cols)
            fill = PatternFill("solid", fgColor="1B5C3F")
            font = Font(bold=True, color="FFFFFF")
            for cell in ws[1]:
                cell.fill = fill
                cell.font = font
            for _, row in df.iterrows():
                ws.append([_clean(v) for v in row])

        # Ensure parent directory exists
        SAVE_PATH.parent.mkdir(parents=True, exist_ok=True)
        wb.save(SAVE_PATH)

        now = datetime.now().strftime("%H:%M:%S")
        st.session_state[_STATUS_KEY] = "ok"
        st.session_state[_TIME_KEY]   = now
        return True

    except Exception as e:
        import streamlit as st
        st.session_state[_STATUS_KEY] = f"error: {e}"
        st.session_state[_TIME_KEY]   = datetime.now().strftime("%H:%M:%S")
        return False


def load_session() -> dict | None:
    """Load the auto-saved file on startup. Returns None if not found."""
    if not SAVE_PATH.exists():
        return None
    try:
        raw = pd.ExcelFile(SAVE_PATH, engine="openpyxl")
        dfs = {}
        for sheet in SCHEMA:
            if sheet in raw.sheet_names:
                df = raw.parse(sheet)
                df.columns = [str(c).strip() for c in df.columns]
                dfs[sheet] = df
        # Also load any extra sheets not in SCHEMA (future-proof)
        for sheet in raw.sheet_names:
            if sheet not in dfs:
                df = raw.parse(sheet)
                df.columns = [str(c).strip() for c in df.columns]
                dfs[sheet] = df
        if not dfs:
            return None
        return dfs
    except Exception:
        return None


def _clean(v):
    """Convert values to Excel-safe types."""
    if v is None:
        return None
    try:
        if isinstance(v, float) and pd.isna(v):
            return None
    except Exception:
        pass
    try:
        if isinstance(v, pd.Timestamp):
            return v.date()
    except Exception:
        pass
    try:
        # Catch pandas NA / NAT
        if pd.isna(v):
            return None
    except Exception:
        pass
    return v
